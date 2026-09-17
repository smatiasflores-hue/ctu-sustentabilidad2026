from io import BytesIO
import re
from docx import Document
import geopandas as gpd
import pandas as pd
import requests
import streamlit as st
from streamlit_folium import st_folium
import folium

# Configuración de la página en modo ancho y reducción del padding superior
st.set_page_config(
    page_title="Certificado Técnico Urbanístico",
    page_icon="📄",
    layout="wide",
)

# Estilos CSS avanzados: Control de espacios y flujo natural para evitar hojas vacías
st.markdown(
    """
    <style>
        /* Reducir el espacio superior predeterminado de Streamlit */
        .block-container {
            padding-top: 1.0rem !important;
            padding-bottom: 2rem !important;
        }

        /* 1. Tipografía Global */
        html, body, [class*="css"] {
            font-family: 'Arial', Helvetica, sans-serif !important;
            color: #2c3e50;
        }

        /* 2. Estilo de la Portada / Título Principal */
        h1 {
            font-size: 26px !important;
            font-weight: 800 !important;
            text-transform: uppercase !important;
            color: #0d3b66 !important;
            margin-top: -5px !important;
            margin-bottom: 2px !important;
        }

        /* 3. Valores de las métricas */
        [data-testid="stMetricValue"] {
            font-size: 20px !important;
            font-weight: 700 !important;
            color: #1f77b4;
            white-space: normal !important;
            overflow: visible !important;
        }
        
        /* 4. Etiquetas de las métricas */
        [data-testid="stMetricLabel"] {
            font-size: 12px !important;
            font-weight: 600 !important;
            color: #555555;
            white-space: normal !important;
            text-transform: uppercase;
        }

        /* 5. Títulos principales (2 al 10): Mayúsculas y Bold */
        h2 {
            font-size: 16px !important;
            font-weight: 700 !important;
            text-transform: uppercase;
            margin-top: 15px !important;
            margin-bottom: 6px !important;
            border-bottom: 2px solid #1f77b4;
            padding-bottom: 2px;
            color: #0d3b66;
            break-after: avoid;
        }

        /* 6. Subtítulos */
        h3 {
            font-size: 14px !important;
            font-weight: 700 !important;
            text-transform: uppercase;
            margin-top: 10px !important;
            margin-bottom: 4px !important;
            color: #333333;
        }

        /* 7. Textos descriptivos */
        p, li, span {
            font-size: 13px !important;
            font-weight: 400 !important;
            line-height: 1.3;
        }

        .contenido-sangria {
            margin-left: 15px;
        }

        /* Forzar formato cuadrado para el mapa */
        iframe {
            width: 100% !important;
            height: 280px !important;
            max-height: 280px !important;
            border-radius: 4px;
        }
        
        /* Configuración estricta para impresión / PDF limpio y fluido */
        @media print {
            .stSidebar { display: none !important; }
            header { display: none !important; }
            button { display: none !important; }
            .stButton { display: none !important; }
            
            .stSelectbox { display: none !important; }
            .stWarning { display: none !important; }
            [data-testid="stDataFrame"] { display: none !important; }
            
            body { background: white; color: black; }
            
            * {
                overflow: visible !important;
                text-overflow: unset !important;
                white-space: normal !important;
            }
            
            div {
                page-break-inside: avoid;
            }
            
            iframe {
                width: 100% !important;
                height: 280px !important;
                max-height: 280px !important;
                page-break-inside: avoid;
            }
        }
    </style>
""",
    unsafe_allow_html=True,
)


# Función para obtener la calle cercana mediante OpenStreetMap (Nominatim)
@st.cache_data(ttl=3600)
def obtener_calle_cercana(lat, lon):
  try:
    url = f"https://nominatim.openstreetmap.org/reverse?format=jsonv2&lat={lat}&lon={lon}"
    headers = {"User-Agent": "CertificadoTecnicoUrbanistico/1.0"}
    response = requests.get(url, headers=headers, timeout=3)
    if response.status_code == 200:
      data = response.json()
      address = data.get("address", {})
      calle = (
          address.get("road")
          or address.get("pedestrian")
          or address.get("suburb")
          or "Calle no identificada"
      )
      return calle
  except Exception:
    pass
  return "No disponible"


# Función auxiliar para extraer el número y letra de parcela desde un valor CCA
def extraer_parcela_de_cca(cca_str):
  try:
    segmento_par = cca_str[35:] if len(cca_str) >= 35 else cca_str
    match_par = re.search(r"(\d+)([A-Za-z]*)", segmento_par)
    if match_par:
      num_p = match_par.group(1).lstrip("0").rstrip("0")
      if not num_p:
        num_p = match_par.group(1).lstrip("0")
      letra_p = match_par.group(2).upper()
      return f"{num_p}{letra_p}" if letra_p else (num_p if num_p else "-")
  except Exception:
    pass
  return "-"


# Función para generar el documento Word basado en una plantilla (.docx)
def generar_documento_word(contexto_datos):
  try:
    doc = Document("plantilla_certificado.docx")
  except Exception:
    doc = Document()
    doc.add_heading("Certificado Técnico Urbanístico", 0)

  reemplazos = {
      "{{PDA}}": str(contexto_datos.get("pda", "")),
      "{{PARTIDO}}": str(contexto_datos.get("partido", "")),
      "{{CIRCUNSCRIPCION}}": str(contexto_datos.get("circunscripcion", "")),
      "{{SECCION}}": str(contexto_datos.get("seccion", "")),
      "{{MANZANA}}": str(contexto_datos.get("manzana", "")),
      "{{PARCELA}}": str(contexto_datos.get("parcela", "")),
      "{{CALLE}}": str(contexto_datos.get("calle", "")),
      "{{ZONA}}": str(contexto_datos.get("zona", "")),
      "{{FOS}}": str(contexto_datos.get("fos", "")),
      "{{FOT}}": str(contexto_datos.get("fot", "")),
      "{{ALTURA}}": str(contexto_datos.get("altura", "")),
      "{{AREA}}": str(contexto_datos.get("area", "")),
  }

  for p in doc.paragraphs:
    for clave, valor in reemplazos.items():
      if clave in p.text:
        p.text = p.text.replace(clave, valor)

  for tabla in doc.tables:
    for fila in tabla.rows:
      for celda in fila.cells:
        for clave, valor in reemplazos.items():
          if clave in celda.text:
            celda.text = celda.text.replace(clave, valor)

  buffer = BytesIO()
  doc.save(buffer)
  buffer.seek(0)
  return buffer


# Encabezado superior izquierdo compacto
st.markdown(
    "<p style='font-size:12px; font-weight:800; color:#333;"
    " text-transform:uppercase; letter-spacing:1px; margin-bottom:0px;'>COMISIÓN"
    " DE SUSTENTABILIDAD - <strong>CAUBAUNO</strong></p>",
    unsafe_allow_html=True,
)
st.title("CERTIFICADO TÉCNICO URBANÍSTICO - LA PLATA")
st.markdown(
    "<p style='font-size:12px; color:#666; margin-top:-2px; margin-bottom:8px;'>"
    "SISTEMA DE CONSULTA Y GESTIÓN DE PARCELAS (MÁS DE 400.000"
    " REGISTROS).</p>",
    unsafe_allow_html=True,
)


# Carga optimizada del archivo CSV desde GitHub Releases
@st.cache_data
def cargar_datos():
  url_csv = "https://github.com/smatiasflores-hue/ctu-sustentabilidad2026/releases/download/v1.0/datos.csv"
  df = pd.read_csv(
      url_csv,
      sep=";",
      encoding="latin-1",
      low_memory=False,
      on_bad_lines="skip",
  )
  return df


# Carga optimizada del archivo GeoJSON de lotes desde GitHub Releases
@st.cache_data
def cargar_geojson():
  url_geojson = "https://github.com/smatiasflores-hue/ctu-sustentabilidad2026/releases/download/v1.0/lotes.geojson"
  return gpd.read_file(url_geojson)


try:
  df = cargar_datos()

  if "busqueda_activa" not in st.session_state:
    st.session_state.busqueda_activa = False
  if "partida_buscada" not in st.session_state:
    st.session_state.partida_buscada = ""

  # ----------------------------------------------------
  # BARRA LATERAL: CONSULTA POR PARTIDA
  # ----------------------------------------------------
  st.sidebar.header("🔍 Consulta por Partida")
  st.sidebar.markdown("**Estructura:** `055` + `[ 6 dígitos de Partida ]`")

  col1, col2 = st.sidebar.columns([1, 2])
  with col1:
    st.text_input("Partido", value="055", disabled=True)
  with col2:
    partida_input = st.text_input(
        "Nº Partida", max_chars=6, placeholder="Ej: 47562"
    )

  consultar = st.sidebar.button("Consultar Parcela")

  if consultar:
    if partida_input:
      st.session_state.busqueda_activa = True
      st.session_state.partida_buscada = partida_input
    else:
      st.sidebar.warning("Por favor ingrese un número de partida.")

  # ----------------------------------------------------
  # LÓGICA DE FILTRADO USANDO LA MEMORIA DE SESIÓN
  # ----------------------------------------------------
  df_filtrado = pd.DataFrame()

  if st.session_state.busqueda_activa and st.session_state.partida_buscada:
    partida_limpia = st.session_state.partida_buscada.strip().zfill(6)
    pda_completo = f"055{partida_limpia}"

    df["PDA_limpio"] = (
        df["PDA"].astype(str).str.split(".").str[0].str.zfill(9)
    )
    df_filtrado = df[df["PDA_limpio"] == pda_completo]

    if not df_filtrado.empty:
      st.sidebar.success(
          f"¡Se encontraron {len(df_filtrado)} registro(s) para {pda_completo}!"
      )

      st.markdown('<div id="seccion-ficha"></div>', unsafe_allow_html=True)

      # Selector de coincidencias múltiples
      if len(df_filtrado) > 1:
        st.warning(
            f"⚠️ Se encontraron {len(df_filtrado)} registros coincidentes para"
            " esta partida. Seleccione cuál desea visualizar en la ficha:"
        )

        opciones = {
            f"Fila {idx} - CCA: {row.get('CCA', 'N/D')} (Zona: {row.get('designacio', 'N/D')})": idx
            for idx, row in df_filtrado.iterrows()
        }

        seleccion_str = st.selectbox(
            "Seleccionar registro a consultar:", list(opciones.keys())
        )
        indice_seleccionado = opciones[seleccion_str]
        row = df_filtrado.loc[indice_seleccionado]
      else:
        row = df_filtrado.iloc[0]

      # Extracción de campos
      cca_val = str(row.get("CCA", ""))

      partido_val = cca_val[0:3] if len(cca_val) >= 3 else "055"
      circunscripcion_val = cca_val[3:5] if len(cca_val) >= 5 else "-"

      seccion_val = "-"
      for char in cca_val[5:12]:
        if char.isalpha():
          seccion_val = char
          break
      if seccion_val == "-" and len(cca_val) >= 7:
        seccion_val = cca_val[6].strip("0") if cca_val[6] != "0" else "-"

      try:
        segmento_mza = cca_val[28:35] if len(cca_val) >= 35 else cca_val
        nums = re.findall(r"\d+", segmento_mza)
        if nums:
          candidatos = [n for n in nums if len(n) >= 3]
          val_bruto = candidatos[0] if candidatos else nums[0]
          manzana_val = val_bruto.lstrip("0").rstrip("0")
          if not manzana_val:
            manzana_val = val_bruto.lstrip("0")
        else:
          manzana_val = "N/D"
      except Exception:
        manzana_val = "N/D"

      parcela_val = extraer_parcela_de_cca(cca_val)

      df_cca_match = df[df["CCA"].astype(str) == cca_val]

      dec_ma_list = (
          df_cca_match["dec_ma"].dropna().astype(str).unique().tolist()
          if "dec_ma" in df.columns
          else [str(row.get("dec_ma", "N/D"))]
      )
      obs2_list = (
          df_cca_match["observacio_2"].dropna().astype(str).unique().tolist()
          if "observacio_2" in df.columns
          else [str(row.get("observacio_2", "N/D"))]
      )

      # ====================================================
      # 1. DATOS
      # ====================================================
      st.header("1. DATOS")

      with st.container():
        st.markdown('<div class="contenido-sangria">', unsafe_allow_html=True)

        col_mapa, col_datos = st.columns([1, 2], gap="large")

        calle_detectada = "Calculando..."
        linderos_vecinos = gpd.GeoDataFrame()

        with col_mapa:
          st.subheader("Ubicación del Lote")
          try:
            gdf = cargar_geojson()
            col_match = None
            for c in ["CCA", "cca", "PDA", "pda", "Partida"]:
              if c in gdf.columns:
                col_match = c
                break

            if col_match:
              gdf_parcela = gdf[gdf[col_match].astype(str) == cca_val]
              if not gdf_parcela.empty:
                if gdf_parcela.crs != "EPSG:4326":
                  gdf_parcela = gdf_parcela.to_crs("EPSG:4326")

                centroid = gdf_parcela.unary_union.centroid
                lat, lon = centroid.y, centroid.x

                calle_detectada = obtener_calle_cercana(lat, lon)

                # Búsqueda espacial de lotes linderos en radio de 25 metros
                try:
                  gdf_metro = gdf_parcela.to_crs(epsg=32721)
                  buffer_metro = gdf_metro.buffer(25)
                  buffer_wgs84 = gpd.GeoSeries(
                      buffer_metro, crs="EPSG:32721"
                  ).to_crs("EPSG:4326")
                  linderos_cercanos = gdf[
                      gdf.to_crs("EPSG:4326").intersects(buffer_wgs84.iloc[0])
                  ]
                  linderos_vecinos = linderos_cercanos[
                      linderos_cercanos[col_match].astype(str) != cca_val
                  ]
                except Exception:
                  linderos_vecinos = gpd.GeoDataFrame()

                m = folium.Map(
                    location=[lat, lon],
                    zoom_start=19,
                    tiles="OpenStreetMap",
                    zoom_control=False,
                    dragging=False,
                    scrollWheelZoom=False,
                )

                if not linderos_vecinos.empty:
                  if linderos_vecinos.crs != "EPSG:4326":
                    linderos_vecinos = linderos_vecinos.to_crs("EPSG:4326")
                  folium.GeoJson(
                      linderos_vecinos,
                      style_function=lambda x: {
                          "fillColor": "#d3d3d3",
                          "color": "#808080",
                          "weight": 1,
                          "fillOpacity": 0.3,
                      },
                      tooltip="Lote Lindero",
                  ).add_to(m)

                folium.GeoJson(
                    gdf_parcela,
                    style_function=lambda x: {
                        "fillColor": "#1f77b4",
                        "color": "#0d3b66",
                        "weight": 2.5,
                        "fillOpacity": 0.7,
                    },
                    tooltip=(
                        f"Parcela Seleccionada | Calle: {calle_detectada} |"
                        f" Manzana: {manzana_val} | Parcela: {parcela_val}"
                    ),
                ).add_to(m)

                st_folium(m, width=320, height=280)
                st.metric(label="Calle Referencia", value=calle_detectada)
              else:
                st.info(
                    "No se encontró un polígono geométrico asociado al CCA"
                    f" `{cca_val}`."
                )
            else:
              st.warning("El archivo `lotes.geojson` no posee columna de enlace.")
          except Exception as map_error:
            st.info(f"Cargue el archivo `lotes.geojson`. (Error: {map_error})")

          # ====================================================
          # SECCIÓN: Listado de Lotes Linderos con Número y Letra de Parcela
          # ====================================================
          st.markdown("<br>", unsafe_allow_html=True)
          st.subheader("Lotes Linderos")
          if not linderos_vecinos.empty:
            col_id = (
                col_match
                if col_match in linderos_vecinos.columns
                else linderos_vecinos.columns[0]
            )
            
            # Recorremos cada lote lindero vecino y le aplicamos la misma extracción de parcela
            linderos_procesados = []
            for _, row_lindero in linderos_vecinos.iterrows():
              cca_lindero = str(row_lindero.get(col_id, ""))
              num_letra_parcela = extraer_parcela_de_cca(cca_lindero)
              if num_letra_parcela and num_letra_parcela != "-":
                linderos_procesados.append(num_letra_parcela)

            # Eliminamos duplicados si los hubiera
            linderos_procesados = sorted(list(set(linderos_procesados)))

            if linderos_procesados:
              for idx, parc in enumerate(linderos_procesados[:10], 1):
                st.markdown(
                    f"<p style='margin: 0px 0px 4px 0px; font-size:12px;'"
                    f" color:#555;'>• Parcela {parc}</p>",
                    unsafe_allow_html=True,
                )
            else:
              st.markdown(
                  "<p style='font-size:12px; color:#666;'>No se pudieron extraer"
                  " las parcelas linderas.</p>",
                  unsafe_allow_html=True,
              )
          else:
            st.markdown(
                "<p style='font-size:12px; color:#666;'>No se detectaron lotes"
                " linderos cercanos.</p>",
                unsafe_allow_html=True,
            )

        with col_datos:
          st.subheader("a. Datos Catastrales")
          c_cat1, c_cat2, c_cat3 = st.columns(3)
          with c_cat1:
            st.metric(label="PDA", value=pda_completo)
            st.metric(label="Sección", value=seccion_val)
          with c_cat2:
            st.metric(label="Partido", value=partido_val)
            st.metric(label="Manzana", value=manzana_val)
          with c_cat3:
            st.metric(label="Circunscripción", value=circunscripcion_val)
            st.metric(label="Parcela", value=parcela_val)

          st.markdown("")

          st.subheader("b. Parámetros Urbanísticos")
          st.metric(
              label="Descripción del Área",
              value=str(row.get("descripcio", "N/D")),
          )
          st.metric(
              label="Descripción Secundaria",
              value=str(row.get("descripcio_2", "N/D")),
          )

          c_urb1, c_urb2, c_urb3, c_urb4 = st.columns(4)
          with c_urb1:
            st.metric(label="Zona", value=str(row.get("designacio", "N/D")))
          with c_urb2:
            st.metric(label="FOS", value=str(row.get("fos", "N/D")))
          with c_urb3:
            st.metric(label="FOT", value=str(row.get("fota", "N/D")))
          with c_urb4:
            st.metric(
                label="Altura Máx.", value=str(row.get("hmax", "N/D"))
            )

          st.subheader("c. Ordenanzas Territoriales")
          if dec_ma_list:
            for item in dec_ma_list:
              st.markdown(f"- {item}")
          else:
            st.markdown("- N/D")

          st.subheader("d. Ordenanzas Municipales")
          if obs2_list:
            for item in obs2_list:
              st.markdown(f"- {item}")
          else:
            st.markdown("- N/D")

        st.markdown("</div>", unsafe_allow_html=True)

      # ====================================================
      # 2 A 10. TÍTULOS PRINCIPALES
      # ====================================================
      st.header("2. ORIENTACIÓN Y VENTILACIÓN / DISEÑO PASIVO")
      with st.container():
        st.markdown('<div class="contenido-sangria">', unsafe_allow_html=True)
        st.write(
            "Análisis de asoleamiento, vientos predominantes y estrategias de"
            " diseño bioclimático pasivo para el lote seleccionado."
        )
        st.markdown("</div>", unsafe_allow_html=True)

      st.header("3. TECNOLOGIA CONSTRUCTIVA")
      with st.container():
        st.markdown('<div class="contenido-sangria">', unsafe_allow_html=True)
        st.write(
            "Sistemas constructivos aptos (Tradicional, Steel Framing, Wood"
            " Framing, paneles SIP) y normativas de aplicación."
        )
        st.markdown("</div>", unsafe_allow_html=True)

      st.header("4. INSTALACIONES")
      with st.container():
        st.markdown('<div class="contenido-sangria">', unsafe_allow_html=True)
        st.write(
            "Requerimientos y factibilidad de servicios sanitarios, eléctricos,"
            " gas y desagües pluviales/cloacales."
        )
        st.markdown("</div>", unsafe_allow_html=True)

      st.header("5. CLIMATIZACIÓN")
      with st.container():
        st.markdown('<div class="contenido-sangria">', unsafe_allow_html=True)
        st.write(
            "Estrategias de acondicionamiento térmico activo y pasivo,"
            " envolvente y eficiencia energética."
        )
        st.markdown("</div>", unsafe_allow_html=True)

      st.header("6. ENERGÍAS ALTERNATIVAS")
      with st.container():
        st.markdown('<div class="contenido-sangria">', unsafe_allow_html=True)
        st.write(
            "Incorporación de sistemas de energías renovables (paneles solares"
            " fotovoltaicos, calentadores de agua solares, etc.)."
        )
        st.markdown("</div>", unsafe_allow_html=True)

      st.header("7. PATOLOGÍAS")
      with st.container():
        st.markdown('<div class="contenido-sangria">', unsafe_allow_html=True)
        st.write(
            "Evaluación de riesgos ambientales, napas freáticas, humedad y"
            " precauciones estructurales del sector."
        )
        st.markdown("</div>", unsafe_allow_html=True)

      st.header("8. ILUMINACIÓN NATURAL")
      with st.container():
        st.markdown('<div class="contenido-sangria">', unsafe_allow_html=True)
        st.write(
            "Factores de iluminación natural, dimensiones mínimas de vanos y"
            " factor de luz diurna según normativa."
        )
        st.markdown("</div>", unsafe_allow_html=True)

      st.header("9. ACCESIBILIDAD")
      with st.container():
        st.markdown('<div class="contenido-sangria">', unsafe_allow_html=True)
        st.write(
            "Criterios de accesibilidad universal, circulaciones horizontales y"
            " verticales, y adecuación a normativas vigentes."
        )
        st.markdown("</div>", unsafe_allow_html=True)

      st.header("10. RESULTADO DIAGNOSTICO")
      with st.container():
        st.markdown('<div class="contenido-sangria">', unsafe_allow_html=True)
        st.success(
            "**Dictamen Urbanístico y Ambiental:** Parcela apta para desarrollo"
            " según parámetros de FOS, FOT y altura máxima establecidos para"
            " la zona."
        )
        st.markdown("</div>", unsafe_allow_html=True)

      # ----------------------------------------------------
      # BOTÓN DE DESCARGA DE WORD OFICIAL (.DOCX)
      # ----------------------------------------------------
      st.markdown("---")
      st.subheader("📥 Generación de Documento Oficial")

      datos_para_docx = {
          "pda": pda_completo,
          "partido": partido_val,
          "circunscripcion": circunscripcion_val,
          "seccion": seccion_val,
          "manzana": manzana_val,
          "parcela": parcela_val,
          "calle": calle_detectada,
          "zona": str(row.get("designacio", "N/D")),
          "fos": str(row.get("fos", "N/D")),
          "fot": str(row.get("fota", "N/D")),
          "altura": str(row.get("hmax", "N/D")),
          "area": str(row.get("descripcio", "N/D")),
      }

      archivo_docx = generar_documento_word(datos_para_docx)

      col_v1, col_btn_dw, col_v2 = st.columns([1, 2, 1])
      with col_btn_dw:
        st.download_button(
            label="📄 Descargar Certificado en Word (.docx)",
            data=archivo_docx,
            file_name=f"Certificado_Urbanistico_{pda_completo}.docx",
            mime=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            use_container_width=True,
            type="primary",
        )

    else:
      st.sidebar.error("No se encontró ninguna parcela con ese número.")
  else:
    st.info(
        "👈 Ingrese los 6 dígitos de la Partida en la barra lateral y presione"
        ' "Consultar Parcela" para ver los datos de la parcela.'
    )

  if not df_filtrado.empty:
    st.dataframe(df_filtrado, use_container_width=True)

except Exception as e:
  st.error(
      f"Ocurrió un error al leer el archivo 'datos.csv'. Detalle técnico: {e}"
  )
